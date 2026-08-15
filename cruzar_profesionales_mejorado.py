#!/usr/bin/env python3
"""
Script mejorado para cruzar profesionales en emisiones otoacústicas.
Incluye detección de días inhábiles y resumen por profesional.

Detecta automáticamente festivos de Chile y marca exámenes en días inhábiles.
"""

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from rapidfuzz import fuzz
import holidays
from datetime import datetime

def normalizar_nombre(nombre):
    """Normaliza un nombre para comparación fuzzy."""
    if pd.isna(nombre):
        return ""
    return str(nombre).upper().strip()

def buscar_profesional_por_nombre(nombre_buscado, nombres_dict, umbral=75):
    """
    Busca un profesional por nombre usando fuzzy matching.
    Retorna (profesional, score) o (None, score) si no encuentra.
    """
    nombre_norm = normalizar_nombre(nombre_buscado)
    if not nombre_norm:
        return None, 0

    mejor_match = None
    mejor_score = 0

    for nombre_registrado, profesional in nombres_dict.items():
        score = fuzz.partial_ratio(nombre_norm, nombre_registrado)
        if score > mejor_score:
            mejor_score = score
            mejor_match = profesional

    if mejor_score >= umbral:
        return mejor_match, mejor_score
    return None, mejor_score

def es_dia_inhabil(fecha, festivos_cl):
    """
    Verifica si una fecha es día inhábil en Chile.
    Inhábil = sábado (5), domingo (6) o festivo.
    """
    if pd.isna(fecha):
        return False

    # Convertir a datetime si es necesario
    if isinstance(fecha, str):
        try:
            fecha = pd.to_datetime(fecha)
        except:
            return False

    # Verificar si es fin de semana
    if fecha.weekday() >= 5:  # 5=sábado, 6=domingo
        return True

    # Verificar si es festivo
    fecha_date = fecha.date() if hasattr(fecha, 'date') else fecha
    return fecha_date in festivos_cl

def procesar_emisiones_mejorado(archivo_datos, archivo_rellenar, archivo_salida):
    """
    Cruza los datos de profesionales con las emisiones otoacústicas.
    Incluye detección de días inhábiles y estadísticas por profesional.

    Retorna un diccionario con:
    - 'df': DataFrame procesado
    - 'resumen_profesionales': Resumen de hábiles/inhabiles por profesional
    - 'examenes_inhabiles': Lista de exámenes en días inhábiles
    """

    print("Leyendo archivo de datos...")
    xl_file = pd.ExcelFile(archivo_datos)
    hojas = xl_file.sheet_names
    print(f"  Hojas encontradas: {hojas}")

    # Inicializar festivos de Chile (todos los años que aparezcan en los datos)
    festivos_cl = holidays.Chile()

    # Leer todas las hojas
    dict_rut_por_hoja = {}
    nombres_todos = {}
    datos_profesionales = {}  # Para guardar fecha de evaluación junto con el profesional

    for hoja in hojas:
        df = pd.read_excel(archivo_datos, sheet_name=hoja)
        df['Rut  Madre'] = df['Rut  Madre'].astype(str).str.strip()

        # Diccionario RUT → (profesional, fecha_evaluacion)
        dict_rut_por_hoja[hoja] = dict(zip(
            df['Rut  Madre'],
            zip(df['Nombre Tecnólogo médico'], df['Fecha de evaluación '])
        ))

        # Diccionario nombre → profesional (para fuzzy matching)
        for n, p in zip(df['Nombre Madre'], df['Nombre Tecnólogo médico']):
            n_norm = normalizar_nombre(n)
            if n_norm and not pd.isna(p):
                nombres_todos[n_norm] = p

    print(f"Datos cargados: {sum(len(d) for d in dict_rut_por_hoja.values())} madres en total")

    # Leer archivo a rellenar
    print("\nLeyendo archivo a rellenar...")
    df_rellenar = pd.read_excel(archivo_rellenar)
    df_rellenar['RUT'] = df_rellenar['RUT'].astype(str).str.strip()
    print(f"  {len(df_rellenar)} filas para procesar")

    # Procesar cada fila
    resultados = []
    metodo_llenado = []
    es_inhabil = []
    examenes_inhabiles_list = []

    print("\nProcesando cruces...")
    for idx, row in df_rellenar.iterrows():
        rut = row['RUT']
        nombre = row['NOMBRE PAC']
        fecha_eval = None
        encontrado = False

        # Intentar por RUT en todas las hojas
        for hoja in hojas:
            if rut in dict_rut_por_hoja[hoja]:
                profesional, fecha_eval = dict_rut_por_hoja[hoja][rut]
                resultados.append(profesional)
                metodo_llenado.append(f'RUT_{hoja}')
                encontrado = True
                break

        if not encontrado:
            # Intentar por nombre (fuzzy matching)
            profesional, score = buscar_profesional_por_nombre(nombre, nombres_todos, umbral=75)
            if profesional:
                resultados.append(profesional)
                metodo_llenado.append(f'NOMBRE({int(score)}%)')
            else:
                resultados.append(None)
                metodo_llenado.append('NO_ENCONTRADO')

        # Verificar si es día inhábil
        inhabil = es_dia_inhabil(fecha_eval, festivos_cl)
        es_inhabil.append(inhabil)

        if inhabil and resultados[idx]:
            examenes_inhabiles_list.append({
                'fecha': fecha_eval,
                'profesional': resultados[idx],
                'madre': nombre,
                'rut': rut
            })

    df_rellenar['PROFESIONAL'] = resultados
    df_rellenar['_ES_INHABIL'] = es_inhabil
    df_rellenar['_METODO'] = metodo_llenado

    # Generar resumen por profesional
    resumen_profesionales = {}
    for prof, es_inh in zip(resultados, es_inhabil):
        if prof:
            if prof not in resumen_profesionales:
                resumen_profesionales[prof] = {'total': 0, 'habil': 0, 'inhabil': 0}
            resumen_profesionales[prof]['total'] += 1
            if es_inh:
                resumen_profesionales[prof]['inhabil'] += 1
            else:
                resumen_profesionales[prof]['habil'] += 1

    # Guardar en Excel
    print(f"\nGuardando en {archivo_salida}...")
    df_salida = df_rellenar.drop(columns=['_METODO', '_ES_INHABIL'])
    df_salida.to_excel(archivo_salida, index=False, sheet_name='Hoja1')

    # Marcar celdas con color
    wb = load_workbook(archivo_salida)
    ws = wb.active

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")

    # Encontrar columnas
    prof_col = None
    for col_idx, cell in enumerate(ws[1], 1):
        if cell.value == 'PROFESIONAL':
            prof_col = col_idx
            break

    if prof_col:
        for row_idx, (metodo, inhabil) in enumerate(zip(metodo_llenado, es_inhabil), start=2):
            cell = ws.cell(row=row_idx, column=prof_col)

            # Prioridad: si es inhábil, amarillo; si es búsqueda por nombre, amarillo; si no encontrado, naranja
            if inhabil:
                cell.fill = yellow_fill
            elif 'NOMBRE' in metodo:
                cell.fill = yellow_fill
            elif 'NO_ENCONTRADO' in metodo:
                cell.fill = orange_fill

    wb.save(archivo_salida)

    # Resumen en consola
    print("\n" + "=" * 60)
    print("RESUMEN DE CRUZAMIENTO")
    print("=" * 60)
    print(f"Total procesadas: {len(df_rellenar)}")
    print(f"Encontradas por RUT: {sum(1 for m in metodo_llenado if 'RUT_' in m)}")
    print(f"Encontradas por nombre: {sum(1 for m in metodo_llenado if 'NOMBRE' in m)}")
    print(f"No encontradas: {metodo_llenado.count('NO_ENCONTRADO')}")
    print(f"Exámenes en días inhábiles: {sum(es_inhabil)}")
    print("=" * 60)

    print("\nRESUMEN POR PROFESIONAL:")
    print("-" * 60)
    for prof in sorted(resumen_profesionales.keys()):
        stats = resumen_profesionales[prof]
        print(f"{prof:30} | Total: {stats['total']:3} | Hábiles: {stats['habil']:3} | Inhábiles: {stats['inhabil']:3}")

    if examenes_inhabiles_list:
        print("\nExámenes en días inhábiles:")
        for exam in examenes_inhabiles_list:
            fecha_str = exam['fecha'].strftime('%Y-%m-%d %A') if hasattr(exam['fecha'], 'strftime') else str(exam['fecha'])
            print(f"  {fecha_str}: {exam['profesional']} - {exam['madre']}")

    print(f"\n✓ Archivo guardado: {archivo_salida}")

    return {
        'df': df_rellenar,
        'resumen_profesionales': resumen_profesionales,
        'examenes_inhabiles': examenes_inhabiles_list,
        'metodo_llenado': metodo_llenado,
        'es_inhabil': es_inhabil
    }

if __name__ == "__main__":
    import sys
    if len(sys.argv) != 4:
        print("Uso: python cruzar_profesionales_mejorado.py <archivo_datos.xlsx> <archivo_rellenar.xlsx> <salida.xlsx>")
        sys.exit(1)

    procesar_emisiones_mejorado(sys.argv[1], sys.argv[2], sys.argv[3])