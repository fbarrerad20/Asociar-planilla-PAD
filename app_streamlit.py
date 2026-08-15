#!/usr/bin/env python3
"""
Aplicación Streamlit para procesar emisiones otoacústicas.
Interfaz web para cruzar profesionales, detectar días inhábiles y enviar reportes.
"""

import streamlit as st
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from rapidfuzz import fuzz
import holidays
from datetime import datetime
import smtplib
from email.mime.multipart import MIMEMultipart
from email.mime.base import MIMEBase
from email.mime.text import MIMEText
from email import encoders
import io
import tempfile
import os

# Configurar la página
st.set_page_config(page_title="Asociar pacientes PAD", layout="wide")
st.title("📊 Asociar pacientes PAD")
st.markdown("Cruza profesionales, detecta días inhábiles y genera reportes automáticos")

# Configuración de correo (en el sidebar)
st.sidebar.header("⚙️ Configuración")
st.sidebar.info("""
**Credenciales de correo:**
Para usar Gmail, necesitas una contraseña de aplicación.
[Instrucciones aquí](https://support.google.com/accounts/answer/185833)
""")

remitente_email = st.sidebar.text_input(
    "Correo remitente (Gmail/Outlook):",
    value="",
    type="password",
    key="email_input"
)
remitente_password = st.sidebar.text_input(
    "Contraseña de aplicación:",
    value="",
    type="password",
    key="password_input"
)
destinatario_email = "f.barrerad20@gmail.com"
st.sidebar.write(f"📧 Reportes se enviarán a: {destinatario_email}")

# Funciones auxiliares
def normalizar_nombre(nombre):
    """Normaliza un nombre para comparación fuzzy."""
    if pd.isna(nombre):
        return ""
    return str(nombre).upper().strip()

def buscar_profesional_por_nombre(nombre_buscado, nombres_dict, umbral=75):
    """Busca un profesional por nombre usando fuzzy matching."""
    nombre_norm = normalizar_nombre(nombre_buscado)
    if not nombre_norm:
        return None, 0

    mejor_match = None
    mejor_score = 0

    for nombre_registrado, profesional in nombres_dict.items():
        score = fuzz.partial_ratio(nombre_norm, nombre_registrado)
        if score > mejor_score:
            mejor_score = score
            mejor_match = profesional

    if mejor_score >= umbral:
        return mejor_match, mejor_score
    return None, mejor_score

def es_dia_inhabil(fecha, festivos_cl):
    """Verifica si una fecha es día inhábil en Chile."""
    if pd.isna(fecha):
        return False

    if isinstance(fecha, str):
        try:
            fecha = pd.to_datetime(fecha)
        except:
            return False

    if fecha.weekday() >= 5:  # sábado=5, domingo=6
        return True

    fecha_date = fecha.date() if hasattr(fecha, 'date') else fecha
    return fecha_date in festivos_cl

def procesar_emisiones(archivo_datos, archivo_rellenar):
    """Procesa el cruce de emisiones y detecta días inhábiles."""

    # Leer archivo de datos
    xl_file = pd.ExcelFile(archivo_datos)
    hojas = xl_file.sheet_names
    festivos_cl = holidays.Chile()

    dict_rut_por_hoja = {}
    nombres_todos = {}

    for hoja in hojas:
        df = pd.read_excel(archivo_datos, sheet_name=hoja)
        df['Rut  Madre'] = df['Rut  Madre'].astype(str).str.strip()

        dict_rut_por_hoja[hoja] = dict(zip(
            df['Rut  Madre'],
            zip(df['Nombre Tecnólogo médico'], df['Fecha de evaluación '])
        ))

        for n, p in zip(df['Nombre Madre'], df['Nombre Tecnólogo médico']):
            n_norm = normalizar_nombre(n)
            if n_norm and not pd.isna(p):
                nombres_todos[n_norm] = p

    # Leer archivo a rellenar
    df_rellenar = pd.read_excel(archivo_rellenar)
    df_rellenar['RUT'] = df_rellenar['RUT'].astype(str).str.strip()

    # Procesar cada fila
    resultados = []
    metodo_llenado = []
    es_inhabil = []
    examenes_inhabiles_list = []

    for idx, row in df_rellenar.iterrows():
        rut = row['RUT']
        nombre = row['NOMBRE PAC']
        fecha_eval = None
        encontrado = False

        for hoja in hojas:
            if rut in dict_rut_por_hoja[hoja]:
                profesional, fecha_eval = dict_rut_por_hoja[hoja][rut]
                resultados.append(profesional)
                metodo_llenado.append(f'RUT_{hoja}')
                encontrado = True
                break

        if not encontrado:
            profesional, score = buscar_profesional_por_nombre(nombre, nombres_todos, umbral=75)
            if profesional:
                resultados.append(profesional)
                metodo_llenado.append(f'NOMBRE({int(score)}%)')
            else:
                resultados.append(None)
                metodo_llenado.append('NO_ENCONTRADO')

        inhabil = es_dia_inhabil(fecha_eval, festivos_cl)
        es_inhabil.append(inhabil)

        if inhabil and resultados[idx]:
            examenes_inhabiles_list.append({
                'fecha': fecha_eval,
                'profesional': resultados[idx],
                'madre': nombre,
                'rut': rut
            })

    df_rellenar['PROFESIONAL'] = resultados
    df_rellenar['_ES_INHABIL'] = es_inhabil
    df_rellenar['_METODO'] = metodo_llenado

    # Generar resumen por profesional
    resumen_profesionales = {}
    for prof, es_inh in zip(resultados, es_inhabil):
        if prof:
            if prof not in resumen_profesionales:
                resumen_profesionales[prof] = {'total': 0, 'habil': 0, 'inhabil': 0}
            resumen_profesionales[prof]['total'] += 1
            if es_inh:
                resumen_profesionales[prof]['inhabil'] += 1
            else:
                resumen_profesionales[prof]['habil'] += 1

    return {
        'df': df_rellenar,
        'resumen_profesionales': resumen_profesionales,
        'examenes_inhabiles': examenes_inhabiles_list,
        'metodo_llenado': metodo_llenado,
        'es_inhabil': es_inhabil
    }

def generar_excel_procesado(resultado, archivo_salida):
    """Genera el Excel con colores y formato."""
    df_salida = resultado['df'].drop(columns=['_METODO', '_ES_INHABIL'])
    df_salida.to_excel(archivo_salida, index=False, sheet_name='Hoja1')

    wb = load_workbook(archivo_salida)
    ws = wb.active

    yellow_fill = PatternFill(start_color="FFFF00", end_color="FFFF00", fill_type="solid")
    orange_fill = PatternFill(start_color="FFA500", end_color="FFA500", fill_type="solid")

    prof_col = None
    for col_idx, cell in enumerate(ws[1], 1):
        if cell.value == 'PROFESIONAL':
            prof_col = col_idx
            break

    if prof_col:
        for row_idx, (metodo, inhabil) in enumerate(zip(resultado['metodo_llenado'], resultado['es_inhabil']), start=2):
            cell = ws.cell(row=row_idx, column=prof_col)
            if inhabil or 'NOMBRE' in metodo:
                cell.fill = yellow_fill
            elif 'NO_ENCONTRADO' in metodo:
                cell.fill = orange_fill

    wb.save(archivo_salida)

def generar_cuerpo_correo(resultado, nombre_archivo):
    """Genera el cuerpo del correo con el resumen."""
    cuerpo = f"""
Reporte de Emisiones Otoacústicas
Generado: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}

Archivo: {nombre_archivo}

ESTADÍSTICAS GENERALES:
─────────────────────────────────────────────────────────
Total de exámenes procesados: {len(resultado['df'])}
Encontrados por RUT: {sum(1 for m in resultado['metodo_llenado'] if 'RUT_' in m)}
Encontrados por nombre: {sum(1 for m in resultado['metodo_llenado'] if 'NOMBRE' in m)}
No encontrados: {resultado['metodo_llenado'].count('NO_ENCONTRADO')}
Exámenes en días inhábiles (doble pago): {sum(resultado['es_inhabil'])}

RESUMEN POR PROFESIONAL:
─────────────────────────────────────────────────────────
"""

    for prof in sorted(resultado['resumen_profesionales'].keys()):
        stats = resultado['resumen_profesionales'][prof]
        cuerpo += f"{prof:30} | Total: {stats['total']:3} | Hábiles: {stats['habil']:3} | Inhábiles: {stats['inhabil']:3}\n"

    if resultado['examenes_inhabiles']:
        cuerpo += "\nDETALLE DE EXÁMENES EN DÍAS INHÁBILES:\n"
        cuerpo += "─────────────────────────────────────────────────────────\n"
        for exam in resultado['examenes_inhabiles']:
            fecha_str = exam['fecha'].strftime('%Y-%m-%d (%A)') if hasattr(exam['fecha'], 'strftime') else str(exam['fecha'])
            cuerpo += f"{fecha_str} | {exam['profesional']} | {exam['madre']}\n"

    cuerpo += "\n\nNotas:\n"
    cuerpo += "🟨 Amarillo en PROFESIONAL = Examen en día inhábil (doble pago) o búsqueda por nombre\n"
    cuerpo += "🟠 Naranja en PROFESIONAL = No encontrado en bases de datos\n"

    return cuerpo

def enviar_correo(remitente, password, destinatario, asunto, cuerpo, archivo_excel):
    """Envía un correo con el archivo adjunto."""
    try:
        # Crear mensaje
        mensaje = MIMEMultipart()
        mensaje['From'] = remitente
        mensaje['To'] = destinatario
        mensaje['Subject'] = asunto

        # Agregar cuerpo
        mensaje.attach(MIMEText(cuerpo, 'plain'))

        # Agregar archivo
        if archivo_excel:
            # Leer el contenido del archivo
            with open(archivo_excel, 'rb') as adjunto:
                parte = MIMEBase('application', 'octet-stream')
                parte.set_payload(adjunto.read())

            encoders.encode_base64(parte)
            parte.add_header('Content-Disposition', 'attachment', filename=os.path.basename(archivo_excel))
            mensaje.attach(parte)

        # Enviar
        if '@gmail.com' in remitente:
            servidor = smtplib.SMTP_SSL('smtp.gmail.com', 465)
        else:  # Outlook/Hotmail
            servidor = smtplib.SMTP_SSL('smtp-mail.outlook.com', 465)

        servidor.login(remitente, password)
        servidor.send_message(mensaje)
        servidor.quit()

        return True, "Correo enviado exitosamente"
    except Exception as e:
        return False, f"Error al enviar correo: {str(e)}"

# Interfaz principal
st.header("📁 Sube tus archivos")

col1, col2 = st.columns(2)

with col1:
    st.subheader("1. Archivo de datos (EOA)")
    archivo_datos = st.file_uploader(
        "Selecciona el Excel con todos los datos de profesionales",
        type=['xlsx'],
        key='datos'
    )

with col2:
    st.subheader("2. Archivo a rellenar (EMISIONES)")
    archivo_rellenar = st.file_uploader(
        "Selecciona el Excel de emisiones a completar",
        type=['xlsx'],
        key='rellenar'
    )

if archivo_datos and archivo_rellenar:
    if st.button("🚀 Procesar archivos", type="primary"):
        try:
            with st.spinner("Procesando... esto puede tardar unos segundos"):
                # Guardar archivos temporales
                with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as f_datos:
                    f_datos.write(archivo_datos.read())
                    ruta_datos = f_datos.name

                with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as f_rellenar:
                    f_rellenar.write(archivo_rellenar.read())
                    ruta_rellenar = f_rellenar.name

                # Procesar
                resultado = procesar_emisiones(ruta_datos, ruta_rellenar)

                # Generar Excel procesado
                with tempfile.NamedTemporaryFile(delete=False, suffix='.xlsx') as f_salida:
                    ruta_salida = f_salida.name

                generar_excel_procesado(resultado, ruta_salida)

                # Mostrar resultados
                st.success("✅ Procesamiento completado")

                # Estadísticas
                col1, col2, col3, col4 = st.columns(4)
                with col1:
                    st.metric("Total procesadas", len(resultado['df']))
                with col2:
                    st.metric("Encontradas por RUT", sum(1 for m in resultado['metodo_llenado'] if 'RUT_' in m))
                with col3:
                    st.metric("Encontradas por nombre", sum(1 for m in resultado['metodo_llenado'] if 'NOMBRE' in m))
                with col4:
                    st.metric("Días inhábiles", sum(resultado['es_inhabil']))

                # Resumen por profesional
                st.subheader("📊 Resumen por Profesional")
                resumen_df = pd.DataFrame([
                    {
                        'Profesional': prof,
                        'Total': stats['total'],
                        'Hábiles': stats['habil'],
                        'Inhábiles': stats['inhabil']
                    }
                    for prof, stats in sorted(resultado['resumen_profesionales'].items())
                ])
                st.dataframe(resumen_df, use_container_width=True)

                # Detalle de días inhábiles
                if resultado['examenes_inhabiles']:
                    st.subheader("⚠️ Exámenes en días inhábiles (doble pago)")
                    inhabiles_df = pd.DataFrame(resultado['examenes_inhabiles'])
                    st.dataframe(inhabiles_df, use_container_width=True)

                # Botón descargar
                st.subheader("📥 Descargar resultado")
                with open(ruta_salida, 'rb') as f:
                    st.download_button(
                        label="Descargar Excel procesado",
                        data=f.read(),
                        file_name="emisiones_procesadas.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                    )

                # Enviar correo
                st.subheader("📧 Enviar reporte")
                if remitente_email and remitente_password:
                    if st.button("Enviar correo a fiscalización"):
                        with st.spinner("Enviando correo..."):
                            asunto = f"Reporte de Emisiones Otoacústicas - {datetime.now().strftime('%Y-%m-%d')}"
                            cuerpo = generar_cuerpo_correo(resultado, archivo_rellenar.name)
                            exito, mensaje = enviar_correo(
                                remitente_email,
                                remitente_password,
                                destinatario_email,
                                asunto,
                                cuerpo,
                                ruta_salida
                            )
                            if exito:
                                st.success(f"✅ {mensaje}")
                            else:
                                st.error(f"❌ {mensaje}")
                else:
                    st.warning("⚠️ Configura tus credenciales de correo en el panel de la izquierda para enviar reportes")

                # Limpiar archivos temporales
                os.unlink(ruta_datos)
                os.unlink(ruta_rellenar)

        except Exception as e:
            st.error(f"❌ Error al procesar: {str(e)}")

else:
    st.info("👆 Sube ambos archivos para comenzar")

st.markdown("---")
st.caption("© 2026 - Sistema de Gestión de Emisiones Otoacústicas")
